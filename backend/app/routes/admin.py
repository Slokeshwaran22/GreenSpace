from io import BytesIO
from typing import Literal, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook, Workbook
from pydantic import BaseModel, Field
from mysql.connector import Error

from app.database import get_connection
from app.dependencies import require_admin
from app.security import (
    normalize_email,
    is_valid_gmail,
    validate_password,
    hash_password,
)


router = APIRouter(prefix="/admin", tags=["Admin"])


# Request model for sending a message.
class SendMessageRequest(BaseModel):
    recipient_id: Optional[int] = Field(default=None, gt=0)
    message_type: Literal["announcement", "general"] = "general"
    title: str = Field(min_length=1, max_length=200)
    body: str = Field(min_length=1, max_length=10000)


@router.get("/admindashboard")
def admin_dashboard(
    admin: dict = Depends(require_admin),
):
    connection = None
    cursor = None

    try:
        connection = get_connection()
        cursor = connection.cursor(dictionary=True)

        # Count all registered accounts.
        cursor.execute("SELECT COUNT(*) AS total FROM users")
        total_users = cursor.fetchone()["total"]

        # Count unique users with unexpired sessions.
        cursor.execute("""
            SELECT COUNT(DISTINCT user_id) AS total
            FROM user_sessions
            WHERE expires_at > NOW()
        """)
        active_users = cursor.fetchone()["total"]

        # Retrieve registered user details.
        cursor.execute("""
            SELECT id, name, email, phone, role
            FROM users
            ORDER BY id ASC
        """)
        users = cursor.fetchall()

        return {
            "total_registered_users": total_users,
            "currently_logged_in": active_users,
            "users": users,
        }

    except Error as error:
        print(f"Admin dashboard error: {error}")
        raise HTTPException(
            status_code=500,
            detail="Unable to load admin dashboard.",
        ) from error

    finally:
        if cursor:
            cursor.close()

        if connection and connection.is_connected():
            connection.close()


@router.get("/messages")
def get_admin_messages(
    admin: dict = Depends(require_admin),
):
    connection = None
    cursor = None

    try:
        connection = get_connection()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("""
            SELECT
                m.id,
                m.sender_id,
                m.recipient_id,
                m.message_type,
                m.title,
                m.body,
                COUNT(DISTINCT mr.user_id) AS recipient_count,
                GROUP_CONCAT(
                    DISTINCT u.name
                    ORDER BY u.name
                    SEPARATOR ', '
                ) AS recipient_names
            FROM messages m
            LEFT JOIN message_recipients mr
                ON mr.message_id = m.id
            LEFT JOIN users u
                ON u.id = mr.user_id
            WHERE m.sender_id = %s
            GROUP BY
                m.id,
                m.sender_id,
                m.recipient_id,
                m.message_type,
                m.title,
                m.body
            ORDER BY m.id DESC
            LIMIT 50
        """, (admin["id"],))

        messages = cursor.fetchall()

        return {"messages": messages}

    except Error as error:
        print(f"Admin messages error: {error}")
        raise HTTPException(
            status_code=500,
            detail="Unable to load sent messages.",
        ) from error

    finally:
        if cursor:
            cursor.close()

        if connection and connection.is_connected():
            connection.close()


@router.post("/messages")
def send_message(
    message: SendMessageRequest,
    admin: dict = Depends(require_admin),
):
    connection = None
    cursor = None

    # Remove accidental leading/trailing whitespace.
    title = message.title.strip()
    body = message.body.strip()

    if not title:
        raise HTTPException(
            status_code=400,
            detail="Message title cannot be empty.",
        )

    if not body:
        raise HTTPException(
            status_code=400,
            detail="Message body cannot be empty.",
        )

    try:
        connection = get_connection()
        cursor = connection.cursor(dictionary=True)

        # Determine recipients.
        if message.recipient_id is None:
            # Broadcast only to regular users.
            cursor.execute("""
                SELECT id
                FROM users
                WHERE LOWER(TRIM(role)) = 'user'
            """)
            recipients = cursor.fetchall()

        else:
            # Check that the selected recipient exists and is a regular user.
            cursor.execute("""
                SELECT id
                FROM users
                WHERE id = %s
                  AND LOWER(TRIM(role)) = 'user'
                LIMIT 1
            """, (message.recipient_id,))

            recipient = cursor.fetchone()

            if not recipient:
                raise HTTPException(
                    status_code=404,
                    detail="Selected user was not found.",
                )

            recipients = [recipient]

        if not recipients:
            raise HTTPException(
                status_code=400,
                detail="No eligible users were found.",
            )

        # Create the message.
        cursor.execute("""
            INSERT INTO messages (
                sender_id,
                recipient_id,
                message_type,
                title,
                body
            )
            VALUES (%s, %s, %s, %s, %s)
        """, (
            admin["id"],
            message.recipient_id,
            message.message_type,
            title,
            body,
        ))

        message_id = cursor.lastrowid

        # Create read-status records for recipients.
        cursor.executemany("""
            INSERT INTO message_recipients (
                message_id,
                user_id
            )
            VALUES (%s, %s)
        """, [
            (message_id, recipient["id"])
            for recipient in recipients
        ])

        connection.commit()

        return {
            "message": "Message sent successfully.",
            "message_id": message_id,
            "message_type": message.message_type,
            "recipient_count": len(recipients),
        }

    except HTTPException:
        if connection and connection.is_connected():
            connection.rollback()
        raise

    except Error as error:
        if connection and connection.is_connected():
            connection.rollback()

        print(f"Admin send message error: {error}")

        raise HTTPException(
            status_code=500,
            detail="Unable to send message.",
        ) from error

    finally:
        if cursor:
            cursor.close()

        if connection and connection.is_connected():
            connection.close()


@router.post("/users/import-excel")
def import_users_from_excel(
    file: UploadFile = File(...),
    admin: dict = Depends(require_admin),
):
    connection = None
    cursor = None
    source_workbook = None

    try:
        # 1. Validate the uploaded file.
        if not file.filename or not file.filename.lower().endswith(".xlsx"):
            raise HTTPException(
                status_code=400,
                detail="Please upload an Excel .xlsx file.",
            )

        file_bytes = file.file.read(5 * 1024 * 1024 + 1)

        if not file_bytes:
            raise HTTPException(
                status_code=400,
                detail="The uploaded Excel file is empty.",
            )

        if len(file_bytes) > 5 * 1024 * 1024:
            raise HTTPException(
                status_code=400,
                detail="Excel file must not exceed 5 MB.",
            )

        # 2. Read the Excel workbook.
        try:
            source_workbook = load_workbook(
                BytesIO(file_bytes),
                read_only=True,
                data_only=True,
            )
        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is not a valid .xlsx workbook.",
            ) from error

        sheet = source_workbook.active

        header_row = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            None,
        )

        if not header_row:
            raise HTTPException(
                status_code=400,
                detail="The Excel file does not contain a header row.",
            )

        # Accept headers with or without a trailing asterisk.
        # Examples: username, username*, email, email*.
        headers = [
            str(value).strip().lower().removesuffix("*").strip()
            if value is not None else ""
            for value in header_row
        ]

        # Username and email are required columns.
        # Phone and password columns are optional.
        required_headers = ["username", "email"]
        optional_headers = ["phone", "password"]

        missing_headers = [
            name for name in required_headers if name not in headers
        ]

        if missing_headers:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Missing required Excel columns: "
                    + ", ".join(missing_headers)
                ),
            )

        # Column order does not matter.
        column_index = {
            name: headers.index(name)
            for name in required_headers + optional_headers
            if name in headers
        }

        # 3. Prepare the result workbook.
        result_workbook = Workbook()
        result_sheet = result_workbook.active
        result_sheet.title = "Import Results"

        result_sheet.append([
            "username",
            "email",
            "phone",
            "creation_status",
            "error_message",
        ])

        total_rows = 0
        created_count = 0
        failed_count = 0

        connection = get_connection()
        cursor = connection.cursor()

        # 4. Process each data row.
        for row in sheet.iter_rows(min_row=2, values_only=True):

            # Skip completely empty rows.
            if not any(
                value is not None and str(value).strip()
                for value in row
            ):
                continue

            total_rows += 1

            def cell_value(column_name):
                index = column_index.get(column_name)

                if index is None or index >= len(row):
                    return None

                return row[index]

            username_value = cell_value("username")
            email_value = cell_value("email")
            phone_value = cell_value("phone")
            password_value = cell_value("password")

            username = (
                str(username_value).strip()
                if username_value is not None else ""
            )

            raw_email = (
                str(email_value).strip()
                if email_value is not None else ""
            )
            email = normalize_email(raw_email) if raw_email else ""

            phone = (
                str(phone_value).strip()
                if phone_value is not None else ""
            )

            password = (
                str(password_value)
                if password_value is not None else ""
            )

            status = "FAILED"
            error_message = ""

            try:
                # Username is required.
                if not username:
                    raise ValueError("Username is required.")

                if len(username) > 60:
                    raise ValueError(
                        "Username must not exceed 60 characters."
                    )

                if not username.replace(" ", "").isalpha():
                    raise ValueError(
                        "Username must contain only letters and spaces."
                    )

                # Email is required.
                if not email:
                    raise ValueError("Email is required.")

                if not is_valid_gmail(email):
                    raise ValueError(
                        "Please enter a valid Gmail address."
                    )

                # Phone is optional.
                if phone:
                    if not phone.isdigit() or len(phone) != 10:
                        raise ValueError(
                            "Phone must contain exactly 10 digits."
                        )
                else:
                    phone = None

                # Password is optional.
                if password:
                    if not 8 <= len(password) <= 16:
                        raise ValueError(
                            "Password must be 8 to 16 characters long."
                        )

                    if not validate_password(password):
                        raise ValueError(
                            "Password must contain uppercase, lowercase, "
                            "a number and a special character."
                        )

                    password_hash = hash_password(password)
                else:
                    password_hash = None

                # Save only the password hash, never plain-text password.
                cursor.execute(
                    """
                    INSERT INTO users (
                        name,
                        phone,
                        email,
                        password_hash,
                        role
                    )
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    (
                        username,
                        phone,
                        email,
                        password_hash,
                        "user",
                    ),
                )

                connection.commit()

                status = "SUCCESS"
                created_count += 1

            except ValueError as error:
                error_message = str(error)
                failed_count += 1

            except Error as error:
                connection.rollback()

                if error.errno == 1062:
                    error_message = (
                        "Email or phone number already exists."
                    )
                else:
                    print(f"Excel import database error: {error}")
                    error_message = (
                        "Database error while creating this user."
                    )

                failed_count += 1

            # Never include the plain-text password in the result workbook.
            result_sheet.append([
                username,
                email,
                phone or "",
                status,
                error_message,
            ])

        if total_rows == 0:
            raise HTTPException(
                status_code=400,
                detail="No user data rows were found in the Excel file.",
            )

        # 5. Add a summary worksheet.
        summary_sheet = result_workbook.create_sheet("Summary")

        summary_sheet.append(["Import Summary", "Count"])
        summary_sheet.append(["Total processed", total_rows])
        summary_sheet.append(["Users created", created_count])
        summary_sheet.append(["Users failed", failed_count])

        # 6. Generate the downloadable Excel result.
        output = BytesIO()
        result_workbook.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    'attachment; filename="GreenSpace_Import_Results.xlsx"'
                ),
                "X-Total-Rows": str(total_rows),
                "X-Users-Created": str(created_count),
                "X-Users-Failed": str(failed_count),
                "Access-Control-Expose-Headers": (
                    "X-Total-Rows, X-Users-Created, X-Users-Failed"
                ),
            },
        )

    except HTTPException:
        raise

    except Exception as error:
        print(f"Excel import error: {error}")
        raise HTTPException(
            status_code=500,
            detail="Unable to process the Excel file.",
        ) from error

    finally:
        if cursor:
            cursor.close()

        if connection and connection.is_connected():
            connection.close()

        if source_workbook:
            source_workbook.close()

        file.file.close()   